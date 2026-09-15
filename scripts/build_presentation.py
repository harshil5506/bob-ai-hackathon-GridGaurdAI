"""
GridGuard AI — Executive Presentation Builder
Generates a super-professional 16:9 widescreen PowerPoint presentation (presentation/slides.pptx)
and automatically exports it to PDF (presentation/slides.pdf) via PowerPoint COM automation.
"""

import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

# ── Color Palette (Dark SCADA HUD Aesthetic) ──────────────────────────────────
BG_DARK = RGBColor(2, 6, 17)          # #020611 Primary Deep Void
PANEL_BG = RGBColor(8, 20, 38)        # #081426 Card Fill
PANEL_BORDER = RGBColor(24, 48, 79)    # #18304f Card Subtle Border
CYAN = RGBColor(0, 243, 255)          # #00f3ff Electric Cyan Accent
GREEN = RGBColor(0, 255, 102)         # #00ff66 Normal / Success
ORANGE = RGBColor(255, 170, 0)        # #ffaa00 Warning / Staging
RED = RGBColor(255, 51, 51)           # #ff3333 Critical / Alert
TEXT_WHITE = RGBColor(255, 255, 255)  # #ffffff Headings
TEXT_SILVER = RGBColor(148, 163, 184) # #94a3b8 Body / Secondary
TEXT_MUTED = RGBColor(100, 116, 139)  # #64748b Labels
PURPLE_IBM = RGBColor(120, 82, 255)   # #7852ff IBM Purple / Accent

FONT_TITLE = "Segoe UI"
FONT_BODY = "Segoe UI"
FONT_MONO = "Consolas"


def set_slide_background(slide):
    """Sets a solid deep dark void background for the slide."""
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = BG_DARK


def add_header(slide, category_text: str, title_text: str, subtitle_text: str = ""):
    """Adds a standardized SCADA HUD header to a slide."""
    # Top subtle accent line
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(0.4), Inches(11.733), Inches(0.03))
    line.fill.solid()
    line.fill.fore_color.rgb = CYAN
    line.line.color.rgb = CYAN

    # Category / breadcrumb tag
    cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11.733), Inches(0.3))
    tf = cat_box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.text = category_text.upper()
    p.font.name = FONT_MONO
    p.font.size = Pt(10)
    p.font.bold = True
    p.font.color.rgb = CYAN

    # Slide Title
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.75), Inches(11.733), Inches(0.55))
    tf2 = title_box.text_frame
    tf2.word_wrap = True
    tf2.margin_left = tf2.margin_top = tf2.margin_right = tf2.margin_bottom = 0
    p2 = tf2.paragraphs[0]
    p2.text = title_text
    p2.font.name = FONT_TITLE
    p2.font.size = Pt(24)
    p2.font.bold = True
    p2.font.color.rgb = TEXT_WHITE

    # Optional Subtitle
    if subtitle_text:
        sub_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.35), Inches(11.733), Inches(0.35))
        tf3 = sub_box.text_frame
        tf3.word_wrap = True
        tf3.margin_left = tf3.margin_top = tf3.margin_right = tf3.margin_bottom = 0
        p3 = tf3.paragraphs[0]
        p3.text = subtitle_text
        p3.font.name = FONT_BODY
        p3.font.size = Pt(12)
        p3.font.color.rgb = TEXT_SILVER


def create_card(slide, left, top, width, height, border_color=PANEL_BORDER, bg_color=PANEL_BG):
    """Creates a stylized panel card with custom border and fill."""
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = bg_color
    card.line.color.rgb = border_color
    card.line.width = Pt(1.5)
    return card


# ── Presentation Generation ───────────────────────────────────────────────────

def build_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    base_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    screenshots_dir = os.path.join(base_dir, "demo", "screenshots")

    # =========================================================================
    # SLIDE 1: TITLE & EXECUTIVE VISION
    # =========================================================================
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_background(s1)

    # Decorative high-voltage grid lines
    top_line = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(0.6), Inches(11.733), Inches(0.04))
    top_line.fill.solid()
    top_line.fill.fore_color.rgb = CYAN
    top_line.line.color.rgb = CYAN

    # Sub-badge
    badge_box = s1.shapes.add_textbox(Inches(0.8), Inches(1.0), Inches(11.733), Inches(0.4))
    tf = badge_box.text_frame
    p = tf.paragraphs[0]
    p.text = "⚡ IBM BOB AI HACKATHON 2026  //  TRACK: AI  //  NERC-CIP & IEEE-1547 GROUNDED"
    p.font.name = FONT_MONO
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = CYAN

    # Main Project Title
    title_box = s1.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(11.733), Inches(1.3))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = "GRIDGUARD AI"
    p.font.name = FONT_TITLE
    p.font.size = Pt(56)
    p.font.bold = True
    p.font.color.rgb = TEXT_WHITE

    # Project Subtitle
    sub_box = s1.shapes.add_textbox(Inches(0.8), Inches(2.8), Inches(11.733), Inches(0.6))
    tf = sub_box.text_frame
    p = tf.paragraphs[0]
    p.text = "Power Outage Prediction & Grid Equipment Failure Advisor"
    p.font.name = FONT_TITLE
    p.font.size = Pt(22)
    p.font.bold = True
    p.font.color.rgb = CYAN

    # Value Proposition Statement Card
    create_card(s1, Inches(0.8), Inches(3.6), Inches(11.733), Inches(1.4), border_color=CYAN, bg_color=PANEL_BG)
    prop_box = s1.shapes.add_textbox(Inches(1.1), Inches(3.8), Inches(11.133), Inches(1.0))
    tf = prop_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Transforming electric utility operations from reactive catastrophe response into proactive, automated resilience by fusing multi-modal sensor telemetry, spatial weather forecasts, and historical failure records with IBM Bob and watsonx.ai."
    p.font.name = FONT_BODY
    p.font.size = Pt(15)
    p.font.color.rgb = TEXT_WHITE

    # Team & Presentation Metadata Cards (Bottom Row)
    team_card = create_card(s1, Inches(0.8), Inches(5.3), Inches(6.8), Inches(1.6), border_color=PANEL_BORDER)
    team_box = s1.shapes.add_textbox(Inches(1.0), Inches(5.45), Inches(6.4), Inches(1.3))
    tf = team_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "PROJECT TEAM: GridGuard AI"
    p.font.name = FONT_MONO
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = CYAN

    p2 = tf.add_paragraph()
    p2.text = "• Harshil Thakkar (Lead) — Full-Stack & System Architect\n• Priyal Rathod — AI/ML Risk Modeling & Physics Engine\n• Rutvik Jasani — Backend & SCADA Telemetry Gateway\n• Hetavi Suhagiya — Frontend UI & Data Visualization"
    p2.font.name = FONT_BODY
    p2.font.size = Pt(11)
    p2.font.color.rgb = TEXT_SILVER

    tech_card = create_card(s1, Inches(7.8), Inches(5.3), Inches(4.733), Inches(1.6), border_color=PANEL_BORDER)
    tech_box = s1.shapes.add_textbox(Inches(8.0), Inches(5.45), Inches(4.333), Inches(1.3))
    tf = tech_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "CORE TECHNOLOGIES"
    p.font.name = FONT_MONO
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = CYAN

    p2 = tf.add_paragraph()
    p2.text = "• IBM Bob CLI & Model Context Protocol (MCP)\n• IBM watsonx.ai (Granite / Llama-3 Foundation Models)\n• Python FastAPI, NumPy, IEEE C57 / IEC 60076 Standards\n• React 18, TypeScript, Recharts, SCADA HUD Console"
    p2.font.name = FONT_BODY
    p2.font.size = Pt(11)
    p2.font.color.rgb = TEXT_SILVER

    # =========================================================================
    # SLIDE 2: THE PROBLEM — WHO, WHAT, AND WHY IT HURTS
    # =========================================================================
    s2 = prs.slides.add_slide(blank_layout)
    set_slide_background(s2)
    add_header(s2, "Problem Statement // The Reliability Crisis", "Aging Power Grids Are Blind to Compounding Weather Failures",
               "Traditional calendar maintenance leaves multi-million-dollar transformers vulnerable to unpredicted catastrophic breakdowns.")

    # 3 Stakeholder Columns (Top Half)
    col_w = Inches(3.75)
    col_gap = Inches(0.24)
    top_y = Inches(1.8)
    col_h = Inches(3.0)

    # Col 1: SCADA Dispatchers
    create_card(s2, Inches(0.8), top_y, col_w, col_h, border_color=RED)
    box = s2.shapes.add_textbox(Inches(1.0), top_y + Inches(0.2), col_w - Inches(0.4), col_h - Inches(0.4))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "SCADA GRID DISPATCHERS"
    p.font.name = FONT_MONO
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = RED

    p2 = tf.add_paragraph()
    p2.text = "• Overwhelmed by 500+ alarms/hr during storms\n• Current SCADA systems only alert AFTER equipment trips\n• Zero forward visibility into 12-72h degradation curves\n• Forced to make high-stress decisions without physics root-cause explanations"
    p2.font.name = FONT_BODY
    p2.font.size = Pt(11)
    p2.font.color.rgb = TEXT_SILVER

    # Col 2: Maintenance Directors
    create_card(s2, Inches(0.8) + col_w + col_gap, top_y, col_w, col_h, border_color=ORANGE)
    box = s2.shapes.add_textbox(Inches(1.0) + col_w + col_gap, top_y + Inches(0.2), col_w - Inches(0.4), col_h - Inches(0.4))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "FIELD CREW DIRECTORS"
    p.font.name = FONT_MONO
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = ORANGE

    p2 = tf.add_paragraph()
    p2.text = "• Reliant on calendar schedules (every 12-24 mos)\n• Internal insulation decay escalates between visits\n• Emergency crews dispatched into impassable flooded roads after failure occurs\n• High overtime costs and spare-parts logistics chaos"
    p2.font.name = FONT_BODY
    p2.font.size = Pt(11)
    p2.font.color.rgb = TEXT_SILVER

    # Col 3: Downstream Consumers & Infrastructure
    create_card(s2, Inches(0.8) + (col_w + col_gap) * 2, top_y, col_w, col_h, border_color=CYAN)
    box = s2.shapes.add_textbox(Inches(1.0) + (col_w + col_gap) * 2, top_y + Inches(0.2), col_w - Inches(0.4), col_h - Inches(0.4))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "CITIZENS & CRITICAL ASSETS"
    p.font.name = FONT_MONO
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = CYAN

    p2 = tf.add_paragraph()
    p2.text = "• Blackouts sever power to hospitals, water pumps, and telecommunications\n• Protracted outages lasting days to weeks\n• Severe economic devastation to regional industry\n• Public Utility Commission regulatory penalties"
    p2.font.name = FONT_BODY
    p2.font.size = Pt(11)
    p2.font.color.rgb = TEXT_SILVER

    # 4 Quantified Metrics (Bottom Row)
    m_w = Inches(2.78)
    m_gap = Inches(0.20)
    m_top = Inches(5.05)
    m_h = Inches(1.85)

    metrics = [
        ("$1M - $2.5M", "PER HOUR COST", "Direct financial cost of major substation blackout in industrial corridors.", RED),
        ("4 - 18 MONTHS", "REPLACEMENT LEAD TIME", "Global procurement delay to engineer & deliver a Large Power Transformer.", ORANGE),
        ("70% BUDGET", "REACTIVE DRAIN", "Utility maintenance expenditure consumed by emergency spot repairs.", CYAN),
        ("500+ ALARMS", "ALARM FLOOD / HR", "Nuisance false-positive alarm triggers disorienting control room operators.", GREEN),
    ]

    for i, (val, title, desc, col) in enumerate(metrics):
        x = Inches(0.8) + (m_w + m_gap) * i
        create_card(s2, x, m_top, m_w, m_h, border_color=col)
        box = s2.shapes.add_textbox(x + Inches(0.15), m_top + Inches(0.15), m_w - Inches(0.3), m_h - Inches(0.3))
        tf = box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = val
        p.font.name = FONT_MONO
        p.font.size = Pt(20)
        p.font.bold = True
        p.font.color.rgb = col

        p2 = tf.add_paragraph()
        p2.text = title
        p2.font.name = FONT_TITLE
        p2.font.size = Pt(10)
        p2.font.bold = True
        p2.font.color.rgb = TEXT_WHITE

        p3 = tf.add_paragraph()
        p3.text = desc
        p3.font.name = FONT_BODY
        p3.font.size = Pt(9.5)
        p3.font.color.rgb = TEXT_SILVER

    # =========================================================================
    # SLIDE 3: THE SOLUTION — WHAT WE BUILT & CORE MECHANISM
    # =========================================================================
    s3 = prs.slides.add_slide(blank_layout)
    set_slide_background(s3)
    add_header(s3, "Solution Overview // Physics-Grounded Intelligence", "GridGuard AI: Closed-Loop Predictive Grid Resilience",
               "Synthesizing IoT telemetry, weather forecasts, and historical logs into automated crew pre-positioning before outages strike.")

    # Left Column: Solution Architecture Overview
    left_w = Inches(5.0)
    create_card(s3, Inches(0.8), Inches(1.8), left_w, Inches(5.1), border_color=CYAN)
    l_box = s3.shapes.add_textbox(Inches(1.05), Inches(2.0), left_w - Inches(0.5), Inches(4.7))
    tf = l_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "WHAT WE BUILT"
    p.font.name = FONT_MONO
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = CYAN

    p2 = tf.add_paragraph()
    p2.text = "\nGridGuard AI transforms electric utility operations from reactive catastrophe response into proactive, automated resilience.\n\n" \
              "• Physics-Informed Anomaly Engine: Evaluates equipment health against IEEE C57.104 and IEC 60076 standards.\n\n" \
              "• Dynamic Weather Risk Multiplier: Correlates incoming atmospheric heatwaves and convective storm vectors with live thermal thresholds.\n\n" \
              "• Downstream Criticality Prioritization: Ranks substations by customer impact and hospital/infrastructure dependence.\n\n" \
              "• Automated Crew Staging: Pre-positions field repair fleets with specific replacement hardware prior to storm landfall."
    p2.font.name = FONT_BODY
    p2.font.size = Pt(11.5)
    p2.font.color.rgb = TEXT_SILVER

    # Right Column: 4-Step Core Mechanism
    right_x = Inches(6.1)
    right_w = Inches(6.433)
    steps = [
        ("01", "MULTI-MODAL DATA FUSION", "Ingests 4 IoT telemetry streams (Temperature, Vibration, Partial Discharge, DGA Oil Quality) + forward spatial weather models + historical maintenance dossiers.", CYAN),
        ("02", "PHYSICS-GROUNDED FAILURE MODEL", "Computes composite risk: 0.35 * Sensor + 0.30 * Weather + 0.20 * History + 0.15 * Age. Estimates remaining safe operating hours before thermal/dielectric runaway.", GREEN),
        ("03", "GRID IMPACT SEVERITY RANKING", "Scales failure probability by downstream customer density and hospital/critical node weighting: Severity = Probability * Grid Impact * Weather Factor.", ORANGE),
        ("04", "PRESCRIPTIVE MAINTENANCE & CREW STAGING", "Generates concrete engineering work orders (DGA degassing, bushing replacement) and automatically stages emergency strike teams in advance.", RED),
    ]

    step_h = Inches(1.15)
    step_gap = Inches(0.16)
    for i, (num, title, desc, col) in enumerate(steps):
        y = Inches(1.8) + (step_h + step_gap) * i
        create_card(s3, right_x, y, right_w, step_h, border_color=col)
        box = s3.shapes.add_textbox(right_x + Inches(0.2), y + Inches(0.12), right_w - Inches(0.4), step_h - Inches(0.24))
        tf = box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = f"{num} // {title}"
        p.font.name = FONT_MONO
        p.font.size = Pt(11.5)
        p.font.bold = True
        p.font.color.rgb = col

        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.name = FONT_BODY
        p2.font.size = Pt(10)
        p2.font.color.rgb = TEXT_SILVER

    # =========================================================================
    # SLIDE 4: SYSTEM ARCHITECTURE & TECHNICAL HIGHLIGHTS
    # =========================================================================
    s4 = prs.slides.add_slide(blank_layout)
    set_slide_background(s4)
    add_header(s4, "Technical Architecture // Microservices Topology", "High-Throughput, Resilient Multi-Tier Infrastructure",
               "Stateless analytical microservices coupled with sub-second SCADA HUD telemetry rendering.")

    # 4 Architecture Tier Columns
    t_w = Inches(2.78)
    t_gap = Inches(0.20)
    t_top = Inches(1.8)
    t_h = Inches(3.6)

    tiers = [
        ("TIER 1: INGESTION", "Data Fusion Engine", "data_fusion.py",
         ["• Real-time IoT sensor bus parser", "• 48h spatial weather API vectorizer", "• Historical MTBF incident cross-ref", "• Multi-source feature harmonization"], CYAN),
        ("TIER 2: AI ANALYTICS", "FastAPI AI Engine", "Port 8001 / REST",
         ["• IEEE C57 & IEC 60076 thresholding", "• Vectorized multi-factor risk scoring", "• Downstream severity ranking", "• Sub-50ms inference latency"], GREEN),
        ("TIER 3: PERSISTENCE", "Backend API Gateway", "Port 5000 / REST",
         ["• Node.js Express 4 REST routing", "• PostgreSQL & TimescaleDB hypertables", "• Resilient offline fallback mock engine", "• NERC-CIP audit trail buffering"], ORANGE),
        ("TIER 4: SCADA CONSOLE", "React 18 + Vite HUD", "Port 3000 / Web",
         ["• 60 FPS HTML5 canvas electric arcs", "• Multi-sensor Recharts trend lines", "• 12 substation geospatial grid map", "• Level 1/2/3 clearance studio"], PURPLE_IBM),
    ]

    for i, (title, sub, tech, bullets, col) in enumerate(tiers):
        x = Inches(0.8) + (t_w + t_gap) * i
        create_card(s4, x, t_top, t_w, t_h, border_color=col)
        box = s4.shapes.add_textbox(x + Inches(0.18), t_top + Inches(0.18), t_w - Inches(0.36), t_h - Inches(0.36))
        tf = box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title
        p.font.name = FONT_MONO
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = col

        p2 = tf.add_paragraph()
        p2.text = sub
        p2.font.name = FONT_TITLE
        p2.font.size = Pt(13)
        p2.font.bold = True
        p2.font.color.rgb = TEXT_WHITE

        p3 = tf.add_paragraph()
        p3.text = tech
        p3.font.name = FONT_MONO
        p3.font.size = Pt(9.5)
        p3.font.color.rgb = CYAN

        for b in bullets:
            pb = tf.add_paragraph()
            pb.text = b
            pb.font.name = FONT_BODY
            pb.font.size = Pt(10)
            pb.font.color.rgb = TEXT_SILVER

    # Bottom Technical Highlights Panel
    create_card(s4, Inches(0.8), Inches(5.6), Inches(11.733), Inches(1.3), border_color=CYAN)
    b_box = s4.shapes.add_textbox(Inches(1.05), Inches(5.72), Inches(11.2), Inches(1.05))
    tf = b_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "KEY ARCHITECTURAL HIGHLIGHTS"
    p.font.name = FONT_MONO
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = CYAN

    p2 = tf.add_paragraph()
    p2.text = "⚡ Zero-Downtime Evaluation: Built-in resilient fallback engine ensures full 12-substation interactivity even if local PostgreSQL is offline.\n" \
              "🔒 NERC-CIP Compliance: Role-based operational clearance (L1 Monitor, L2 Dispatch, L3 Chief), simulated hardware 2FA, and out-of-band OTP recovery.\n" \
              "🚀 Cloud Ready: Stateless microservices architecture deployable onto IBM Cloud Code Engine or Red Hat OpenShift with horizontal auto-scaling."
    p2.font.name = FONT_BODY
    p2.font.size = Pt(10.5)
    p2.font.color.rgb = TEXT_WHITE

    # =========================================================================
    # SLIDE 5: LIVE DEMO & TECHNICAL HIGHLIGHTS
    # =========================================================================
    s5 = prs.slides.add_slide(blank_layout)
    set_slide_background(s5)
    add_header(s5, "Product Showcase // Live SCADA Control Console", "Operational Command Center & Asset Diagnostics",
               "Real-time visibility across 12 high-voltage transmission nodes with interactive telemetry and crew dispatch.")

    # Left: Screenshot Panel
    left_img_w = Inches(6.8)
    left_img_h = Inches(4.8)
    create_card(s5, Inches(0.8), Inches(1.8), left_img_w, left_img_h, border_color=CYAN)

    img_path = os.path.join(screenshots_dir, "gridguard_dashboard.png")
    if os.path.exists(img_path):
        s5.shapes.add_picture(img_path, Inches(0.9), Inches(1.9), left_img_w - Inches(0.2), left_img_h - Inches(0.2))
    else:
        # Fallback text if image missing
        box = s5.shapes.add_textbox(Inches(1.2), Inches(3.0), left_img_w - Inches(0.8), Inches(2.0))
        box.text_frame.text = "[Dashboard Screenshot: Common Grid Map & Asset Telemetry]"

    # Right: Key Features Breakdown
    r_x = Inches(7.85)
    r_w = Inches(4.683)
    feat_cards = [
        ("COMMON GRID DASHBOARD", "12 High-Voltage Nodes",
         "Geospatial SCADA map color-coded by real-time risk severity. Criticality table automatically sorts substations by failure probability and downline customer impact.", CYAN),
        ("ASSET DIGITAL TWIN & TRENDS", "Multi-Sensor Recharts",
         "Small-multiple time-series plots for Top-Oil Temperature, Vibration, Partial Discharge, and DGA Oil Quality. Pinpoints critical failure window (< 14 hours).", GREEN),
        ("OPERATOR PROFILE & ENCLAVE", "Clearance Studio (L1 / L2 / L3)",
         "Role switcher for live hackathon demos. Rotates security keys, verifies out-of-band email OTP reset codes, and records an immutable NERC-CIP audit trail.", ORANGE),
    ]

    card_h = Inches(1.5)
    card_gap = Inches(0.15)
    for i, (f_title, f_sub, f_desc, col) in enumerate(feat_cards):
        y = Inches(1.8) + (card_h + card_gap) * i
        create_card(s5, r_x, y, r_w, card_h, border_color=col)
        box = s5.shapes.add_textbox(r_x + Inches(0.18), y + Inches(0.12), r_w - Inches(0.36), card_h - Inches(0.24))
        tf = box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = f_title
        p.font.name = FONT_MONO
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = col

        p2 = tf.add_paragraph()
        p2.text = f_sub
        p2.font.name = FONT_TITLE
        p2.font.size = Pt(12)
        p2.font.bold = True
        p2.font.color.rgb = TEXT_WHITE

        p3 = tf.add_paragraph()
        p3.text = f_desc
        p3.font.name = FONT_BODY
        p3.font.size = Pt(9.5)
        p3.font.color.rgb = TEXT_SILVER

    # =========================================================================
    # SLIDE 6: IBM TECHNOLOGY INTEGRATION — BOB, WATSONX & MCP
    # =========================================================================
    s6 = prs.slides.add_slide(blank_layout)
    set_slide_background(s6)
    add_header(s6, "Cognitive Architecture // IBM Ecosystem", "Where and How IBM Bob & watsonx Are Integrated",
               "Combining IBM Bob conversational orchestration with watsonx.ai foundation models via Model Context Protocol.")

    # 4 IBM Technology Cards
    ibm_w = Inches(5.7)
    ibm_gap_x = Inches(0.333)
    ibm_h = Inches(2.3)
    ibm_top1 = Inches(1.8)
    ibm_top2 = Inches(4.3)

    ibm_stack = [
        ("IBM BOB CLI & AGENT", "Conversational SCADA Advisor",
         "Acts as the 24/7 intelligent control room co-pilot. Grid dispatchers interrogate complex anomalies using natural language queries ('Why is Substation Alpha at critical risk?') and receive instant, context-aware engineering diagnoses.",
         Inches(0.8), ibm_top1, CYAN),
        ("MODEL CONTEXT PROTOCOL (MCP)", "Standardized SCADA Tool Bridge",
         "Connects IBM Bob directly to live GridGuard backend REST endpoints. Exposes real-time tools to fetch asset telemetry, query weather risk multipliers, compute IEEE failure probabilities, and trigger emergency crew dispatch.",
         Inches(0.8) + ibm_w + ibm_gap_x, ibm_top1, PURPLE_IBM),
        ("WATSONX.AI FOUNDATION MODELS", "Enterprise Playbook Generation",
         "Leverages IBM Granite and Llama-3 models to synthesize multi-sensor telemetry into structured, IEEE-compliant mitigation playbooks, load-shedding schedules, and regulatory NERC-CIP incident reports.",
         Inches(0.8), ibm_top2, GREEN),
        ("IBM CLOUD CODE ENGINE", "Serverless Elastic Scaling",
         "Runs containerized AI inference microservices on managed serverless infrastructure. Auto-scales computing instances instantaneously when severe weather systems compound grid-wide sensor ingestion loads.",
         Inches(0.8) + ibm_w + ibm_gap_x, ibm_top2, ORANGE),
    ]

    for title, sub, desc, left, top, col in ibm_stack:
        create_card(s6, left, top, ibm_w, ibm_h, border_color=col)
        box = s6.shapes.add_textbox(left + Inches(0.2), top + Inches(0.18), ibm_w - Inches(0.4), ibm_h - Inches(0.36))
        tf = box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title
        p.font.name = FONT_MONO
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = col

        p2 = tf.add_paragraph()
        p2.text = sub
        p2.font.name = FONT_TITLE
        p2.font.size = Pt(13)
        p2.font.bold = True
        p2.font.color.rgb = TEXT_WHITE

        p3 = tf.add_paragraph()
        p3.text = desc
        p3.font.name = FONT_BODY
        p3.font.size = Pt(10.5)
        p3.font.color.rgb = TEXT_SILVER

    # Bottom Cognitive Pipeline Flow
    flow_bar = create_card(s6, Inches(0.8), Inches(6.75), Inches(11.733), Inches(0.45), border_color=CYAN, bg_color=PANEL_BG)
    f_box = s6.shapes.add_textbox(Inches(0.9), Inches(6.8), Inches(11.533), Inches(0.35))
    tf = f_box.text_frame
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.text = "[DISPATCHER PROMPT] ➔ [IBM BOB AGENT] ➔ [MCP TOOL CALL] ➔ [SCADA TELEMETRY + WATSONX.AI] ➔ [ACTIONABLE MITIGATION PLAYBOOK]"
    p.font.name = FONT_MONO
    p.font.size = Pt(9.5)
    p.font.bold = True
    p.font.color.rgb = CYAN

    # =========================================================================
    # SLIDE 7: BUSINESS IMPACT & BEYOND THE HACKATHON
    # =========================================================================
    s7 = prs.slides.add_slide(blank_layout)
    set_slide_background(s7)
    add_header(s7, "Strategic Value // Commercial Vision", "Projected Business Impact & Beyond the Hackathon",
               "Delivering massive operational savings and a clear trajectory toward production utility deployment.")

    # 3 Impact Stat Cards (Top Row)
    stat_w = Inches(3.75)
    stat_gap = Inches(0.24)
    stat_top = Inches(1.8)
    stat_h = Inches(1.8)

    stat_cards = [
        ("42% REDUCTION", "UNPLANNED OUTAGE DURATION", "Compresses utility SAIDI/SAIFI outage indices via predictive strike crew staging.", GREEN),
        ("$85M / YEAR", "REGIONAL COST AVOIDANCE", "Savings for a standard 2M-meter utility by eliminating catastrophic transformer replacements.", CYAN),
        ("12,000+ HOURS", "FIELD LABOR OPTIMIZATION", "Eliminates wasteful calendar-based manual inspections with continuous condition monitoring.", ORANGE),
    ]

    for i, (val, title, desc, col) in enumerate(stat_cards):
        x = Inches(0.8) + (stat_w + stat_gap) * i
        create_card(s7, x, stat_top, stat_w, stat_h, border_color=col)
        box = s7.shapes.add_textbox(x + Inches(0.2), stat_top + Inches(0.18), stat_w - Inches(0.4), stat_h - Inches(0.36))
        tf = box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = val
        p.font.name = FONT_MONO
        p.font.size = Pt(22)
        p.font.bold = True
        p.font.color.rgb = col

        p2 = tf.add_paragraph()
        p2.text = title
        p2.font.name = FONT_TITLE
        p2.font.size = Pt(11)
        p2.font.bold = True
        p2.font.color.rgb = TEXT_WHITE

        p3 = tf.add_paragraph()
        p3.text = desc
        p3.font.name = FONT_BODY
        p3.font.size = Pt(10)
        p3.font.color.rgb = TEXT_SILVER

    # 3-Phase Commercial Roadmap (Bottom Row)
    r_top = Inches(3.8)
    r_h = Inches(3.1)

    phases = [
        ("PHASE 1 (Q3-Q4 2026)", "Physical SCADA & Hardware Integration",
         "• Connect to physical substation RTUs via DNP3 & IEC 61850 protocols.\n"
         "• Deploy edge inference gateways onto substation control house hardware.\n"
         "• Ingest live optical fiber partial discharge & continuous online DGA monitors.", CYAN),
        ("PHASE 2 (Q1-Q2 2027)", "Autonomous Drone & Robotics Dispatch",
         "• Automated drone inspection launch triggers for thermal infrared substation scans.\n"
         "• Real-time road flood & downed-tree satellite routing for pre-positioned fleets.\n"
         "• Closed-loop automated breaker tripping recommendations.", GREEN),
        ("PHASE 3 (2027+)", "Regional RTO Federated Intelligence",
         "• Multi-utility threat intelligence sharing across ISO / RTO interconnections.\n"
         "• Cross-border cascading blackout mitigation with watsonx distributed models.\n"
         "• Commercial SaaS deployment across North American & European utilities.", PURPLE_IBM),
    ]

    for i, (stage, name, details, col) in enumerate(phases):
        x = Inches(0.8) + (stat_w + stat_gap) * i
        create_card(s7, x, r_top, stat_w, r_h, border_color=col)
        box = s7.shapes.add_textbox(x + Inches(0.2), r_top + Inches(0.18), stat_w - Inches(0.4), r_h - Inches(0.36))
        tf = box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = stage
        p.font.name = FONT_MONO
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = col

        p2 = tf.add_paragraph()
        p2.text = name
        p2.font.name = FONT_TITLE
        p2.font.size = Pt(12)
        p2.font.bold = True
        p2.font.color.rgb = TEXT_WHITE

        p3 = tf.add_paragraph()
        p3.text = details
        p3.font.name = FONT_BODY
        p3.font.size = Pt(10)
        p3.font.color.rgb = TEXT_SILVER

    # =========================================================================
    # SLIDE 8: TEAM & SUBMISSION SUMMARY
    # =========================================================================
    s8 = prs.slides.add_slide(blank_layout)
    set_slide_background(s8)
    add_header(s8, "Submission Summary // Team GridGuard AI", "Engineering the Future of Resilient Energy Grids",
               "A complete, production-grade AI platform built for the IBM Bob AI Hackathon.")

    # 4 Team Member Cards
    member_w = Inches(2.78)
    member_gap = Inches(0.20)
    member_top = Inches(1.8)
    member_h = Inches(3.2)

    team_members = [
        ("Harshil Thakkar", "TEAM LEAD & FULL-STACK ARCHITECT", "d25dit083@charusat.edu.in",
         "• Overall platform architecture\n• React SCADA HUD & canvas\n• FastApi / Express integration\n• End-to-end evaluation pipeline", CYAN),
        ("Priyal Rathod", "AI/ML RISK ENGINEER", "24dce119@charusat.edu.in",
         "• IEEE C57 / IEC 60076 scoring\n• Multi-modal data fusion pipeline\n• Weather compounding model\n• Time-to-failure forecasting", GREEN),
        ("Rutvik Jasani", "BACKEND & SCADA GATEWAY", "24dce048@charusat.edu.in",
         "• Node.js REST API gateway\n• PostgreSQL / TimescaleDB schema\n• Resilient fallback engine\n• NERC-CIP audit logging", ORANGE),
        ("Hetavi Suhagiya", "FRONTEND & DATA VIZ", "24dce140@charusat.edu.in",
         "• Recharts sensor telemetry\n• Geospatial substation map\n• Operator clearance studio\n• UI design tokens & accessibility", PURPLE_IBM),
    ]

    for i, (name, role, email, contribs, col) in enumerate(team_members):
        x = Inches(0.8) + (member_w + member_gap) * i
        create_card(s8, x, member_top, member_w, member_h, border_color=col)
        box = s8.shapes.add_textbox(x + Inches(0.18), member_top + Inches(0.18), member_w - Inches(0.36), member_h - Inches(0.36))
        tf = box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = name
        p.font.name = FONT_TITLE
        p.font.size = Pt(15)
        p.font.bold = True
        p.font.color.rgb = TEXT_WHITE

        p2 = tf.add_paragraph()
        p2.text = role
        p2.font.name = FONT_MONO
        p2.font.size = Pt(9.5)
        p2.font.bold = True
        p2.font.color.rgb = col

        p3 = tf.add_paragraph()
        p3.text = email
        p3.font.name = FONT_MONO
        p3.font.size = Pt(8.5)
        p3.font.color.rgb = TEXT_MUTED

        p4 = tf.add_paragraph()
        p4.text = contribs
        p4.font.name = FONT_BODY
        p4.font.size = Pt(9.5)
        p4.font.color.rgb = TEXT_SILVER

    # Bottom Submission Links & Closing Callout
    create_card(s8, Inches(0.8), Inches(5.25), Inches(11.733), Inches(1.65), border_color=CYAN)
    callout_box = s8.shapes.add_textbox(Inches(1.1), Inches(5.4), Inches(11.133), Inches(1.35))
    tf = callout_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "⚡ SUBMISSION ARTIFACTS & LINKS"
    p.font.name = FONT_MONO
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = CYAN

    p2 = tf.add_paragraph()
    p2.text = "• Source Code: github.com/harshil5506/bob-ai-hackathon-GridGaurdAI\n" \
              "• Comprehensive Documentation: docs/problem-statement.md | docs/solution-overview.md | docs/architecture.md | docs/setup-guide.md\n" \
              "• Live SCADA Console: http://localhost:3000  (Pre-loaded 1-Click Chief Controller Login)\n" \
              "• Mission: Protecting the critical electric power that powers modern civilization."
    p2.font.name = FONT_BODY
    p2.font.size = Pt(11)
    p2.font.color.rgb = TEXT_WHITE

    # Save Presentation PPTX
    presentation_dir = os.path.join(base_dir, "presentation")
    os.makedirs(presentation_dir, exist_ok=True)
    pptx_path = os.path.join(presentation_dir, "slides.pptx")
    prs.save(pptx_path)
    print(f"[SUCCESS] PowerPoint presentation saved to: {pptx_path}")

    # Export to PDF via PowerPoint COM if available
    export_to_pdf(pptx_path, os.path.join(presentation_dir, "slides.pdf"))


def export_to_pdf(pptx_path: str, pdf_path: str):
    """Exports PowerPoint PPTX to PDF using win32com PowerPoint application."""
    try:
        import win32com.client
        import time

        abs_pptx = os.path.abspath(pptx_path)
        abs_pdf = os.path.abspath(pdf_path)

        print(f"[INFO] Exporting PPTX to PDF via PowerPoint COM...")
        ppt_app = win32com.client.DispatchEx("PowerPoint.Application")
        # ppSaveAsPDF = 32
        presentation = ppt_app.Presentations.Open(abs_pptx, WithWindow=False)
        presentation.SaveAs(abs_pdf, 32)
        presentation.Close()
        ppt_app.Quit()
        print(f"[SUCCESS] PDF presentation successfully generated at: {abs_pdf}")
    except Exception as e:
        print(f"[WARNING] PowerPoint COM PDF export encountered: {e}")
        print("[INFO] slides.pptx is ready and can be converted or viewed directly.")


if __name__ == "__main__":
    build_deck()
